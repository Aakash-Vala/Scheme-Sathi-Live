# -*- coding: utf-8 -*-
import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = pptx.Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Institutional Color Palette
C_NAVY_DEEP = RGBColor(11, 17, 32)      # #0B1120
C_NAVY_DARK = RGBColor(15, 23, 42)      # #0F172A
C_EMERALD_DARK = RGBColor(4, 120, 87)   # #047857
C_EMERALD_MED = RGBColor(5, 150, 105)   # #059669
C_EMERALD_NEON = RGBColor(16, 185, 129) # #10B981
C_EMERALD_BG = RGBColor(236, 253, 245)  # #ECFDF5
C_AMBER = RGBColor(217, 119, 6)         # #D97706
C_AMBER_BG = RGBColor(254, 243, 199)    # #FEF3C7
C_SLATE_DARK = RGBColor(30, 41, 59)     # #1E293B
C_SLATE_TEXT = RGBColor(51, 65, 85)     # #334155
C_MUTED = RGBColor(100, 116, 139)       # #64748B
C_LIGHT_BG = RGBColor(248, 250, 252)    # #F8FAFC
C_WHITE = RGBColor(255, 255, 255)
C_BORDER = RGBColor(226, 232, 240)      # #E2E8F0
C_BORDER_EMERALD = RGBColor(167, 243, 208)
C_BLUE = RGBColor(37, 99, 235)
C_BLUE_BG = RGBColor(239, 246, 255)

LOGO_PATH = os.path.abspath('static/scheme_sathi_logo.png')

def add_header_banner(slide, title_text, slide_num):
    header_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.92))
    header_bg.fill.solid()
    header_bg.fill.fore_color.rgb = C_NAVY_DEEP
    header_bg.line.fill.background()

    # Tricolor Strip
    saffron = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(4.444), Inches(0.04))
    saffron.fill.solid()
    saffron.fill.fore_color.rgb = RGBColor(255, 153, 51)
    saffron.line.fill.background()
    white = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.444), Inches(0), Inches(4.444), Inches(0.04))
    white.fill.solid()
    white.fill.fore_color.rgb = C_WHITE
    white.line.fill.background()
    green = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.888), Inches(0), Inches(4.445), Inches(0.04))
    green.fill.solid()
    green.fill.fore_color.rgb = RGBColor(19, 136, 8)
    green.line.fill.background()

    # Team Pill
    team_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(0.16), Inches(1.85), Inches(0.58))
    team_box.fill.solid()
    team_box.fill.fore_color.rgb = C_NAVY_DARK
    team_box.line.color.rgb = C_AMBER
    team_box.line.width = Pt(1.5)
    tf = team_box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = 'SCHEME SATHI'
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.alignment = PP_ALIGN.CENTER

    # Title
    title_box = slide.shapes.add_textbox(Inches(2.4), Inches(0.12), Inches(7.5), Inches(0.65))
    tf2 = title_box.text_frame
    tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
    p2 = tf2.paragraphs[0]
    p2.text = title_text
    p2.font.size = Pt(17)
    p2.font.bold = True
    p2.font.color.rgb = C_WHITE

    # SIH Badge
    sih_box = slide.shapes.add_textbox(Inches(10.1), Inches(0.10), Inches(2.8), Inches(0.70))
    tf3 = sih_box.text_frame
    tf3.vertical_anchor = MSO_ANCHOR.MIDDLE
    p3 = tf3.paragraphs[0]
    p3.text = 'SMART INDIA HACKATHON 2026'
    p3.font.size = Pt(9.5)
    p3.font.bold = True
    p3.font.color.rgb = C_EMERALD_NEON
    p3.alignment = PP_ALIGN.RIGHT
    p3_sub = tf3.add_paragraph()
    p3_sub.text = 'PS ID: 26092 | MoSJE & NSFDC'
    p3_sub.font.size = Pt(8.5)
    p3_sub.font.color.rgb = C_AMBER
    p3_sub.alignment = PP_ALIGN.RIGHT

    # Footer
    footer_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.18), Inches(13.333), Inches(0.32))
    footer_bg.fill.solid()
    footer_bg.fill.fore_color.rgb = C_NAVY_DEEP
    footer_bg.line.fill.background()

    footer_text = slide.shapes.add_textbox(Inches(0.4), Inches(7.18), Inches(12.5), Inches(0.32))
    ftf = footer_text.text_frame
    ftf.vertical_anchor = MSO_ANCHOR.MIDDLE
    fp = ftf.paragraphs[0]
    fp.text = f'@SIH Idea submission- Template {slide_num}  |  Scheme Sathi - Problem Statement ID: 26092  |  GitHub: https://github.com/Aakash-Vala/Scheme-Sathi  |  Prototype: http://localhost:8000'
    fp.font.size = Pt(7.8)
    fp.font.color.rgb = RGBColor(148, 163, 184)
